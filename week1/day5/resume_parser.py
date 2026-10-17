import os
from pathlib import Path
from groq import Groq
from dotenv import load_dotenv
from pydantic import BaseModel 
from pypdf import PdfReader
from docx import Document
import time

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API KEY is not available")

client = Groq(api_key=my_api_key)

job_description = """
About the job
AI Software Architect | AI-Native Software Development
12 month rolling contract
We’re working with a well-established organisation to hire a very senior Software Engineer / Architect for a role focused on the next generation of AI-enabled software development.
This isn’t a traditional AI/ML role.
You won’t be building AI products – you’ll be building software using AI.
We’re looking for someone with a strong software engineering background who is already using AI coding agents and assistants as a meaningful part of how they design, develop, test and deliver software.
What you'll do
Build and evolve software with AI tools embedded throughout the development lifecycle.
Use AI coding agents such as Claude Code, GitHub Copilot or similar to accelerate development.
Use AI to explore architecture, solve problems, debug and work across unfamiliar codebases.
Apply AI throughout the SDLC, rather than treating it as a standalone technology.
Continue to work hands-on with software while helping shape how AI can improve engineering productivity.
Challenge and critically assess AI-generated output, taking ownership of the quality of what reaches production.
Work as a senior technical contributor, bringing strong engineering judgement and architectural thinking to complex software problems.

What we're looking for
10–15+ years of software engineering experience with a strong hands-on development background.
Experience operating at Principal Engineer, Software Architect, Technical Architect, Chief Architect or similar level.
Strong understanding of modern software development and the full SDLC.
Hands-on experience using AI coding agents, GitHub Copilot, Claude Code or similar tools.
Evidence that AI is genuinely integrated into your software engineering workflow, rather than simply being an area of interest.
Strong technical and architectural judgement, with the ability to assess and challenge AI-generated solutions.
A builder mentality – someone who is still close to the code and enjoys figuring out how technology can be used to build better software.
Previous industry experience is preferred over a purely consulting-based background.

What this role isn'tThis isn't about building LLMs, developing machine learning models or creating AI products.The focus is AI-powered software development – using the latest AI tools to build features, solve engineering problems and make the software development lifecycle faster and more effective.
Why this role?AI is changing the way software gets built.The organisation is looking for someone who is already experimenting with that change – someone who can talk about how they actually use AI agents in their day-to-day engineering work, not just someone who has AI on their CV.If you're a senior software engineer or architect who is deeply hands-on with AI-assisted development and interested in helping shape what modern software engineering looks like, we'd like to talk.
Apply Today!At Hadfield Green, we are an equal opportunities employer and actively welcome applicants from all backgrounds.We’re not a one-size-fits-all type of organisation, so if you need an accommodation or adjustment in advance of meeting due to a disability or impairment, just let us know and we’d be delighted to assist.
"""
model = "openai/gpt-oss-120b"

class JobD(BaseModel):
    role:str
    required_skills:list[str]
    preferred_skills:list[str]
    minimum_experience:float | None
    education_requirements:list[str]
    responsibilities:list[str]

jobd_schema = JobD.model_json_schema()


system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job description and extract structured information from them.

Return ONLY valid JSON matching this schema:
{jobd_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_prompt = f'''
Analyze the following job description:

{job_description}
'''

message_system = {
    "role": "system",
    "content": system_prompt
}

message = {
    "role": "user",
    "content": user_prompt
}

response_format = {
    "type": "json_object"
}

messages = [message_system, message]
response = client.chat.completions.create(
    model=model, 
    messages=messages,
    response_format=response_format
)

answer = response.choices[0].message.content

import json
job_data = json.loads(answer)
job = JobD(**job_data)




class MatchResult(BaseModel):
    score: float
    details: dict

class Experience(BaseModel):
    company:str | None=None
    role:str | None=None
    duration:str | None=None
    description:str | None=None
    skills_used: list[str] = []

class Resume(BaseModel):
    name:str | None=None
    email:str | None=None
    phone: str | None=None

    total_experience_years:float | None=None

    skills: list[str]=[]
    experience: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()

def parse_resume(resume_text):
    system_prompt = f"""
        You are an Expert resume parser.

        Extract information from the resume based on its meaning,
        not only based on exact section headings.

        Different resumes may use different headings.

        for example:
        - Experience
        - Professional Experience
        - Work History
        - Employment
        - Internships

        These may all contain relevant experience.

        Skills may also appear in the skills section, work experience, internships or projects.

        Return ONLY valid JSON matching this schema:

        {resume_schema}

        Important rules:

        1. Do not invent information.
        2. If a value is not available, return null.
        3. If a list has no information, return an empty list.
        4. Include internships inside experiences.
        5. Extract skills mentioned across the entire resume.
    """

    user_prompt = f"""
        Parse the following resume:

        {resume_text}
    """

    message_system = {
        "role": "system",
        "content": system_prompt
    }

    message = {
        "role": "user",
        "content": user_prompt
    }

    response_format = {
        "type": "json_object"
    }

    messages = [message_system, message]

    response = client.chat.completions.create(
        model=model, 
        messages=messages,
        response_format=response_format
    )

    answer = response.choices[0].message.content
    resume_data = json.loads(answer)
    resume = Resume(**resume_data)
    return resume


def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
        You are a HR recruiter.

        Compare the candidate's resume with the job description.

        JOB DESCRIPTION:
        {job.model_dump_json(indent=2)}

        CANDIDATE RESUME:
        {resume.model_dump_json(indent=2)}

        Return JSON matching this schema:

        {match_schema}

        Give me:

        1. Candidate name
        2. Matching skills
        3. Missing important skills 
        4. Whether experience requirement is met
        5. Overall match percentage from 0 to 100
        6. A short final verdict

        Keep the response concise and easy to read.
    """
    message = {
        "role": "user",
        "content": prompt
    }

    response_format = {
        "type": "json_object"
    }

    messages = [message]
    response = client.chat.completions.create(
        model=model, 
        messages=messages,
        response_format=response_format
    )

    answer = response.choices[0].message.content
    data = json.loads(answer)
    return MatchResult(**data)

def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if(page_text):
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if(paragraph.text.strip()):
            text += paragraph.text + "\n"
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if(cell.text.strip()):
                    text += cell.text + "\n"
    return text

def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == "docx":
        return read_docx(file_path)
    else:
        return None

resume_folder = Path("resumes")
all_results=[]

for file_path in resume_folder.iterdir():
    if(file_path.suffix.lower() not in [".pdf", ".docx"]):
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume=parse_resume(resume_text)
    time.sleep(5)
    result = final_score(job, parsed_resume)
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
    all_results.sort(key=lambda candidate: candidate['score'], reverse=True)
    top_2 = all_results[:2]
    worst_2 = all_results[-2:]

    print("TOP 2 CANDIDATES")
    for candidate in top_2:
        print(
            candidate['name'],
            "-",
            candidate['score'],
            "%"
        )
        print(candidate['details'])

    print("LOWEST 2 CANDIDATES")
    for candidate in worst_2:
        print(
            candidate['name'],
            "-",
            candidate['score'],
            "%"
        )
        print(candidate['details'])

